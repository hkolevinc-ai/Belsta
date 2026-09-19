#!/usr/bin/env python3
"""Create a Temu upload workbook from selected BELSTA Shopify variants."""

from __future__ import annotations

import argparse
import concurrent.futures
import csv
import json
import logging
import re
import time
import zipfile
from collections import defaultdict
from dataclasses import dataclass
from html.parser import HTMLParser
from pathlib import Path
from typing import Any
from urllib.parse import parse_qs, urlencode, urlparse
from urllib.request import Request, urlopen
from xml.etree import ElementTree as ET

from openpyxl import load_workbook
from openpyxl.utils import column_index_from_string


SITE = "https://www.belsta.bg"
USER_AGENT = "Mozilla/5.0 (compatible; BELSTA-Temu-Scraper/1.0)"

QUANTITY = 100
PACKAGE_LENGTH_CM = 20
PACKAGE_WIDTH_CM = 10
PACKAGE_HEIGHT_CM = 5
RETAIL_PAIR_WEIGHT_G = 500
HANDLING_TIME = "1 Day"
FULFILLMENT_CHANNEL = "I will ship this item myself"
COUNTRY_OF_ORIGIN = "Ukraine"

# Fields intentionally left blank at the user's request.
SHIPPING_TEMPLATE = ""
MANUFACTURER = ""
EU_RESPONSIBLE_PERSON = ""

CATEGORY_WOMEN_SLIPPERS = "29194"
CATEGORY_MEN_SLIPPERS = "30533"
CATEGORY_GIRLS_SLIPPERS = "29728"
CATEGORY_BOYS_SLIPPERS = "30862"

FOOT_LENGTH_CM = {
    "23": 15.0,
    "24": 15.5,
    "25": 16.0,
    "26": 17.0,
    "27": 18.0,
    "30": 19.5,
    "31": 20.0,
    "32": 21.0,
    "33": 21.5,
    "34": 22.2,
    "35": 23.0,
    "36": 23.5,
    "37": 24.1,
    "38": 24.8,
    "39": 25.5,
    "40": 26.0,
    # Size 41 differs for women and men and is handled separately.
    "42": 27.6,
    "43": 28.2,
    "44": 28.8,
    "45": 29.5,
    "46": 31.0,
}

COLOR_MAP = {
    "бял": "White",
    "бяла": "White",
    "бяло": "White",
    "white": "White",
    "червен": "Red",
    "червена": "Red",
    "червено": "Red",
    "red": "Red",
    "черен": "Black",
    "черна": "Black",
    "черно": "Black",
    "black": "Black",
    "бежов": "Beige",
    "бежова": "Beige",
    "бежово": "Beige",
    "beige": "Beige",
    "розов": "Pink",
    "розова": "Pink",
    "розово": "Pink",
    "pink": "Pink",
    "пудра": "Dusty rose",
    "powder": "Dusty rose",
    "корал": "Coral",
    "coral": "Coral",
    "бордо": "Burgundy",
    "bordoo": "Burgundy",
    "burgundy": "Burgundy",
    "жълт": "Yellow",
    "жълта": "Yellow",
    "жълто": "Yellow",
    "yellow": "Yellow",
    "сив": "Grey",
    "сива": "Grey",
    "сиво": "Grey",
    "grey": "Grey",
    "gray": "Grey",
    "тъмносив": "Dark Grey",
    "тъмно сив": "Dark Grey",
    "dark grey": "Dark Grey",
    "светлосив": "Light Grey",
    "светло сив": "Light Grey",
    "light grey": "Light Grey",
    "син": "Blue",
    "синя": "Blue",
    "синьо": "Blue",
    "blue": "Blue",
    "тъмносин": "Navy Blue",
    "тъмно син": "Navy Blue",
    "navy": "Navy Blue",
    "navy blue": "Navy Blue",
    "светлосин": "Light Blue",
    "светло син": "Light Blue",
    "light blue": "Light Blue",
    "зелен": "Green",
    "зелена": "Green",
    "зелено": "Green",
    "green": "Green",
    "лилав": "Purple",
    "лилава": "Purple",
    "лилаво": "Purple",
    "purple": "Purple",
    "кафяв": "Brown",
    "кафява": "Brown",
    "кафяво": "Brown",
    "brown": "Brown",
    "оранжев": "Orange",
    "оранжева": "Orange",
    "оранжево": "Orange",
    "orange": "Orange",
    "флорален": "Multicolor",
    "цветен": "Multicolor",
    "многоцветен": "Multicolor",
    "multicolor": "Multicolor",
    "mix": "Multicolor",
    "бронзов": "Brown",
    "бронзова": "Brown",
    "bronze": "Brown",
    "травяной": "Green",
    "златен": "Golden",
    "златна": "Golden",
    "gold": "Golden",
    "golden": "Golden",
    "turquoise": "Cyan",
    "тюркоаз": "Cyan",
    "default": "Black",
    "crimson": "Carmine",
    "milk": "Creamy White",
    "млечен": "Creamy White",
    "сребърен": "Silvery Grey",
    "сребърна": "Silvery Grey",
    "silver": "Silvery Grey",
}

# Visually reviewed BELSTA images that show a wholesale bundle containing many
# pairs. They must never be used in a retail Temu listing.
WHOLESALE_IMAGE_MARKERS = (
    "84b23fa3b996d126681d2788ff34ba9c_",
    "firefly_seed367966",
)
WHOLESALE_IMAGE_FILENAMES = {
    "1_9d7ea284-df25-41b7-9465-177f872d8948.png",
    "1_51684791-009f-4c11-9d76-741ed70bcc42.png",
    "43ae7dc775fdd74636c6ad8863a0196a.jpg",
    "1_0904ac37-49f0-4202-8d8d-cae8db1e258c.png",
    "1_0a8ed8a1-69fa-4ed2-8cea-25e3f1aeebbe.png",
    "1_1da5a3f0-028a-4b2a-8370-0bb9e0dbd57a.png",
}

COLOR_COLLISION_ALTERNATIVES = {
    "Pink": ["Deep Pink", "Bright Pink", "Fuchsia", "Magenta", "Dusty rose"],
    "Blue": ["Royal Blue", "Navy Blue", "Light Blue", "Sky Blue", "Cyan"],
    "Red": ["Deep red", "Scarlet", "Carmine", "Rose Red"],
    "Green": ["Dark Green", "Emerald Green", "Light Green", "Olive Green"],
    "Grey": ["Dark Grey", "Light Grey", "Graphite Color", "Slate Grey"],
    "Brown": ["Dark Brown", "Light Brown", "Coffee", "Chocolate"],
    "Purple": ["Deep purple", "Violet", "Lavender", "Fuchsia"],
}


@dataclass(frozen=True)
class InputSelection:
    row_number: int
    original_url: str
    handle: str
    variant_id: str | None


@dataclass
class OutputVariant:
    product: dict[str, Any]
    variant: dict[str, Any]
    canonical_url: str
    category: str
    size: str
    source_size: str
    source_pack_pairs: int
    retail_from_pack: bool
    source_color: str
    temu_color: str
    sku_images: list[str]
    detail_images: list[str]


class ProductHTMLParser(HTMLParser):
    def __init__(self) -> None:
        super().__init__(convert_charrefs=True)
        self.all_text: list[str] = []
        self.bullets: list[str] = []
        self._current_li: list[str] | None = None
        self._ignored_depth = 0

    def handle_starttag(self, tag: str, attrs: list[tuple[str, str | None]]) -> None:
        if tag in {"script", "style"}:
            self._ignored_depth += 1
        elif tag == "li" and self._ignored_depth == 0:
            self._current_li = []

    def handle_endtag(self, tag: str) -> None:
        if tag in {"script", "style"} and self._ignored_depth:
            self._ignored_depth -= 1
        elif tag == "li" and self._current_li is not None:
            value = normalize_space(" ".join(self._current_li))
            if value and value not in self.bullets:
                self.bullets.append(value)
            self._current_li = None

    def handle_data(self, data: str) -> None:
        if self._ignored_depth:
            return
        value = normalize_space(data)
        if not value:
            return
        self.all_text.append(value)
        if self._current_li is not None:
            self._current_li.append(value)


def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--available", default="input/available.xlsx")
    parser.add_argument("--template", default="input/TEMU_TEMPLATE.xlsx")
    parser.add_argument("--output", default="output/BELSTA_TEMU_UPLOAD.xlsx")
    parser.add_argument("--report", default="output/BELSTA_TEMU_REPORT.csv")
    parser.add_argument("--summary", default="output/BELSTA_TEMU_SUMMARY.txt")
    return parser.parse_args()


def normalize_space(value: Any) -> str:
    return re.sub(r"\s+", " ", str(value or "")).strip()


def parse_product_link(raw_url: str, row_number: int) -> InputSelection:
    parsed = urlparse(raw_url.strip())
    parts = [part for part in parsed.path.split("/") if part]
    if len(parts) < 2 or parts[-2] != "products":
        raise ValueError(f"Row {row_number}: not a BELSTA product URL: {raw_url}")
    handle = parts[-1]
    variant_id = parse_qs(parsed.query).get("variant", [None])[0]
    return InputSelection(row_number, raw_url.strip(), handle, variant_id)


def read_selections(path: Path) -> list[InputSelection]:
    workbook = load_workbook(path, read_only=True, data_only=True)
    sheet = workbook.active
    if sheet.max_column == 1 and sheet.max_row == 1:
        sheet.reset_dimensions()

    selections: list[InputSelection] = []
    for row_number, row in enumerate(sheet.iter_rows(min_row=2, values_only=True), start=2):
        raw = row[0] if row else None
        if raw is None or not str(raw).strip():
            continue
        selections.append(parse_product_link(str(raw), row_number))
    if not selections:
        raise ValueError("available.xlsx does not contain product links in column A.")
    return selections


def fetch_product(handle: str, attempts: int = 4) -> dict[str, Any]:
    url = f"{SITE}/products/{handle}.js"
    last_error: Exception | None = None
    for attempt in range(1, attempts + 1):
        try:
            request = Request(
                url,
                headers={"User-Agent": USER_AGENT, "Accept": "application/json"},
            )
            with urlopen(request, timeout=45) as response:
                payload = json.load(response)
            if not payload.get("variants"):
                raise ValueError("product contains no variants")
            return payload
        except Exception as exc:  # noqa: BLE001 - retry network and parse failures.
            last_error = exc
            if attempt < attempts:
                time.sleep(attempt * 2)
    raise RuntimeError(f"Could not fetch {url}: {last_error}")


def fetch_products(handles: list[str]) -> dict[str, dict[str, Any]]:
    products: dict[str, dict[str, Any]] = {}
    with concurrent.futures.ThreadPoolExecutor(max_workers=6) as executor:
        future_to_handle = {executor.submit(fetch_product, handle): handle for handle in handles}
        for future in concurrent.futures.as_completed(future_to_handle):
            handle = future_to_handle[future]
            products[handle] = future.result()
            logging.info("Fetched %s", handle)
    return products


def option_names(product: dict[str, Any]) -> list[str]:
    names: list[str] = []
    for item in product.get("options", []):
        names.append(normalize_space(item.get("name") if isinstance(item, dict) else item))
    return names


def option_indexes(product: dict[str, Any]) -> tuple[int | None, int | None]:
    size_index: int | None = None
    color_index: int | None = None
    for index, name in enumerate(option_names(product), start=1):
        lowered = name.casefold()
        if color_index is None and any(token in lowered for token in ("color", "colour", "цвят")):
            color_index = index
        if size_index is None and any(token in lowered for token in ("size", "размер")):
            size_index = index
    return size_index, color_index


def variant_option(variant: dict[str, Any], index: int | None) -> str:
    if index is None:
        return ""
    return normalize_space(variant.get(f"option{index}"))


def canonical_image_url(raw: str | None) -> str:
    if not raw:
        return ""
    value = str(raw)
    if value.startswith("//"):
        return "https:" + value
    return value


def is_wholesale_image_url(raw: str | None) -> bool:
    if not raw:
        return False
    filename = urlparse(str(raw)).path.rsplit("/", 1)[-1].casefold()
    return filename in WHOLESALE_IMAGE_FILENAMES or any(
        marker in filename for marker in WHOLESALE_IMAGE_MARKERS
    )


def image_groups_by_color(product: dict[str, Any], color_index: int | None) -> dict[str, list[str]]:
    ordered_images = [
        canonical_image_url(url)
        for url in product.get("images", [])
        if url and not is_wholesale_image_url(url)
    ]
    if not ordered_images:
        return {}

    color_starts: list[tuple[int, str]] = []
    seen: set[tuple[str, str]] = set()
    for variant in product.get("variants", []):
        color = variant_option(variant, color_index) or "Default"
        featured = canonical_image_url((variant.get("featured_image") or {}).get("src"))
        if not featured or is_wholesale_image_url(featured) or (color, featured) in seen:
            continue
        seen.add((color, featured))
        try:
            position = ordered_images.index(featured)
        except ValueError:
            position = len(ordered_images)
            ordered_images.append(featured)
        color_starts.append((position, color))

    if not color_starts:
        return {"Default": ordered_images[:10]}

    color_starts.sort()
    groups: dict[str, list[str]] = defaultdict(list)
    for index, (start, color) in enumerate(color_starts):
        end = color_starts[index + 1][0] if index + 1 < len(color_starts) else len(ordered_images)
        group = ordered_images[start:end] or [ordered_images[start]]
        for image_url in group:
            if image_url and image_url not in groups[color]:
                groups[color].append(image_url)
    return groups


def translate_color(raw_color: str) -> str:
    cleaned = normalize_space(raw_color)
    lowered = cleaned.casefold()
    if lowered in COLOR_MAP:
        return COLOR_MAP[lowered]
    for source, target in COLOR_MAP.items():
        if re.search(rf"(?<!\w){re.escape(source)}(?!\w)", lowered):
            return target
    return "Multicolor"


def temu_color(raw_color: str, category: str) -> str:
    translated = translate_color(raw_color)
    # The Women / Shoes / Slippers category does not offer a Multicolor value.
    # Floral variants are predominantly pink-toned in this catalogue.
    if category == CATEGORY_WOMEN_SLIPPERS and translated == "Multicolor":
        return "Pink"
    return translated


def assign_temu_colors(source_colors: list[str], category: str) -> dict[str, str]:
    """Assign a unique valid Temu color to every distinct source color."""
    assigned: dict[str, str] = {}
    used: set[str] = set()
    for source_color in source_colors:
        if source_color in assigned:
            continue
        base = temu_color(source_color, category)
        candidates: list[str] = [base]
        if translate_color(source_color) == "Multicolor" and category == CATEGORY_WOMEN_SLIPPERS:
            candidates.extend(["Deep Pink", "Bright Pink", "Fuchsia", "Magenta"])
        candidates.extend(COLOR_COLLISION_ALTERNATIVES.get(base, []))
        chosen = next((candidate for candidate in candidates if candidate not in used), None)
        if chosen is None:
            raise ValueError(
                f"Could not assign a unique Temu color for {source_color!r} in category {category}."
            )
        assigned[source_color] = chosen
        used.add(chosen)
    return assigned


def classify_category(product: dict[str, Any]) -> str:
    text = f"{product.get('title', '')} {product.get('handle', '')}".casefold()
    if any(token in text for token in ("myazhki", "mazhki", "мъжки")):
        return CATEGORY_MEN_SLIPPERS
    if any(token in text for token in ("junosheski", "yunosheski", "юношески", "тийн")):
        return CATEGORY_BOYS_SLIPPERS
    if any(token in text for token in ("detski", "детски")):
        # The template has no unisex children's slippers category. Most selected
        # BELSTA children's designs are catalogued under Girls / Slippers.
        return CATEGORY_GIRLS_SLIPPERS
    return CATEGORY_WOMEN_SLIPPERS


def select_variants(
    selections: list[InputSelection], products: dict[str, dict[str, Any]]
) -> tuple[list[OutputVariant], list[str]]:
    selections_by_handle: dict[str, list[InputSelection]] = defaultdict(list)
    for selection in selections:
        selections_by_handle[selection.handle].append(selection)

    output: list[OutputVariant] = []
    warnings: list[str] = []

    for handle in sorted(selections_by_handle):
        product = products[handle]
        size_index, color_index = option_indexes(product)
        variants_by_id = {str(item.get("id")): item for item in product.get("variants", [])}
        chosen_colors: set[str] = set()

        for selection in selections_by_handle[handle]:
            if selection.variant_id is None:
                selected_variant = next(
                    (item for item in product.get("variants", []) if item.get("available")),
                    None,
                )
                if selected_variant is None:
                    warnings.append(
                        f"Row {selection.row_number}: no available default variant was found for {handle}."
                    )
                    continue
                chosen_colors.add(variant_option(selected_variant, color_index) or "Default")
                continue
            selected_variant = variants_by_id.get(selection.variant_id)
            if selected_variant is None:
                warnings.append(
                    f"Row {selection.row_number}: variant {selection.variant_id} was not found for {handle}."
                )
                continue
            chosen_colors.add(variant_option(selected_variant, color_index) or "Default")

        if not chosen_colors:
            warnings.append(f"No valid selected variant remained for {handle}.")
            continue

        groups = image_groups_by_color(product, color_index)
        all_product_images = [
            canonical_image_url(url)
            for url in product.get("images", [])
            if url and not is_wholesale_image_url(url)
        ]
        selected_for_product: list[dict[str, Any]] = []

        for variant in product.get("variants", []):
            if not variant.get("available"):
                continue
            color = variant_option(variant, color_index) or "Default"
            if color not in chosen_colors:
                continue
            selected_for_product.append(variant)

        selected_colors_in_output = {
            variant_option(item, color_index) or "Default" for item in selected_for_product
        }
        detail_images: list[str] = []
        for color in sorted(selected_colors_in_output):
            for image_url in groups.get(color, []):
                if image_url and image_url not in detail_images:
                    detail_images.append(image_url)
        if not detail_images:
            detail_images = all_product_images

        category = classify_category(product)
        ordered_source_colors = list(
            dict.fromkeys(
                variant_option(item, color_index) or "Default" for item in selected_for_product
            )
        )
        temu_colors_by_source = assign_temu_colors(ordered_source_colors, category)

        for variant in selected_for_product:
            source_size = variant_option(variant, size_index) or "One Size"
            retail_sizes = expand_size_range(source_size)
            retail_from_pack = len(retail_sizes) > 1
            source_pack_pairs = count_pairs(source_size) if retail_from_pack else 1
            color = variant_option(variant, color_index) or "Default"
            sku_images = list(groups.get(color, []))
            featured = canonical_image_url((variant.get("featured_image") or {}).get("src"))
            if featured and not is_wholesale_image_url(featured) and featured not in sku_images:
                sku_images.insert(0, featured)
            if not sku_images:
                sku_images = detail_images[:1]

            canonical_url = f"{SITE}/products/{handle}?{urlencode({'variant': variant.get('id')})}"
            for retail_size in retail_sizes:
                output.append(
                    OutputVariant(
                        product=product,
                        variant=variant,
                        canonical_url=canonical_url,
                        category=category,
                        size=retail_size,
                        source_size=source_size,
                        source_pack_pairs=source_pack_pairs,
                        retail_from_pack=retail_from_pack,
                        source_color=color,
                        temu_color=temu_colors_by_source[color],
                        sku_images=sku_images[:10],
                        detail_images=detail_images[:30],
                    )
                )

        if not selected_for_product:
            warnings.append(f"No currently available sizes were found for the selected color(s) of {handle}.")

    # Several wholesale variants can describe different pack compositions for
    # the same product, colour and retail size. Temu needs one row per unique
    # sellable combination, so keep the lowest per-pair source price.
    unique: dict[tuple[str, str, str], OutputVariant] = {}
    for item in output:
        key = (
            str(item.product.get("id")),
            item.source_color.casefold(),
            item.size.casefold(),
        )
        previous = unique.get(key)
        if previous is None or retail_unit_price(item) < retail_unit_price(previous):
            unique[key] = item
    return list(unique.values()), warnings


def html_to_text(raw_html: str) -> str:
    parser = ProductHTMLParser()
    parser.feed(raw_html or "")
    return normalize_space(" ".join(parser.all_text))


def extract_bullets(raw_html: str) -> list[str]:
    parser = ProductHTMLParser()
    parser.feed(raw_html or "")
    return [value[:700] for value in parser.bullets[:6]]


def infer_materials(product: dict[str, Any]) -> tuple[str, str, str]:
    text = html_to_text(product.get("description", "")).casefold()

    if any(token in text for token in ("еко кожа", "еко-кожа", "полиуретанов")):
        upper = "PU"
    elif "деним" in text or "джинс" in text:
        upper = "Denim"
    elif "вълна" in text or "вълнен" in text:
        upper = "Wool"
    else:
        upper = "Fabric"

    if "каучук" in text or "гума" in text:
        sole = "Rubber"
    elif "pvc" in text or "пвц" in text:
        sole = "PVC"
    else:
        sole = "PU"

    if "стелка" in text and ("вълна" in text or "вълнен" in text):
        insole = "Wool"
    elif "стелка" in text and ("еко кожа" in text or "еко-кожа" in text):
        insole = "PU"
    else:
        insole = "Fabric"
    return upper, sole, insole


def infer_pattern(product: dict[str, Any]) -> str:
    text = f"{product.get('title', '')} {html_to_text(product.get('description', ''))}".casefold()
    if "пейсли" in text:
        return "Paisley"
    if "точк" in text:
        return "Polka Dot"
    if "флорал" in text or "цветя" in text or "цветен" in text:
        return "Flowers"
    if "сърце" in text:
        return "Heart"
    if any(token in text for token in ("кот", "мишка", "фламинго", "акула", "плодове")):
        return "Animal Print" if not "плодове" in text else "Fruit"
    if any(token in text for token in ("принт", "бродерия", "апликация")):
        return "Graphic"
    return "Solid color"


def infer_popular_element(product: dict[str, Any]) -> str:
    text = f"{product.get('title', '')} {html_to_text(product.get('description', ''))}".casefold()
    if "пандел" in text:
        return "Bow"
    if "брод" in text:
        return "Embroidery"
    if "цвет" in text:
        return "Flowers"
    return ""


def infer_season(product: dict[str, Any]) -> str:
    text = f"{product.get('title', '')} {html_to_text(product.get('description', ''))}".casefold()
    if any(token in text for token in ("топл", "пух", "вълн", "зимен", "студения сезон")):
        return "Fall/Winter"
    return "All-season"


def infer_closure(product: dict[str, Any]) -> str:
    text = f"{product.get('title', '')} {html_to_text(product.get('description', ''))}".casefold()
    if "велкро" in text or "лепенка" in text:
        return "Hook and Loop Fastener"
    if "ластик" in text:
        return "Elastic Band"
    return "Slip On"


def foot_length(size: str, category: str) -> float | str:
    clean = normalize_space(size)
    if not re.fullmatch(r"\d{2}", clean):
        return ""
    if clean == "41":
        return 27.0 if category == CATEGORY_MEN_SLIPPERS else 26.5
    return FOOT_LENGTH_CM.get(clean, "")


def expand_size_range(size_label: str) -> list[str]:
    """Expand the first numeric shoe-size range into unique retail sizes."""
    clean = normalize_space(size_label)
    match = re.search(r"(?<!\d)(\d{2})\s*[-–—]\s*(\d{2})(?!\d)", clean)
    if not match:
        return [clean or "One Size"]
    start, end = int(match.group(1)), int(match.group(2))
    if start > end or end - start > 20:
        return [clean]
    return [str(size) for size in range(start, end + 1)]


def count_pairs(size_label: str) -> int:
    match = re.search(r"(\d{2})\s*[-–—]\s*(\d{2})", size_label)
    if not match:
        return 1
    start, end = int(match.group(1)), int(match.group(2))
    base = max(1, end - start + 1)
    extras = 0
    for multiplicity in re.findall(r"(\d+)\s*\*\s*\d+", size_label):
        extras += max(0, int(multiplicity) - 1)
    return base + extras


def retail_unit_price(item: OutputVariant) -> float:
    divisor = item.source_pack_pairs if item.retail_from_pack else 1
    return float(item.variant.get("price") or 0) / max(1, divisor)


def retail_product_title(title: str) -> str:
    value = re.sub(r"^\s*Опаковка\s*:\s*", "", normalize_space(title), flags=re.IGNORECASE)
    value = re.sub(r"\s+на\s+едро\b", "", value, flags=re.IGNORECASE)
    return normalize_space(value)


def sku_size_token(size: str) -> str:
    token = re.sub(r"[^A-Za-z0-9]+", "-", normalize_space(size)).strip("-")
    return token or "ONE-SIZE"


def build_workbook_rows(items: list[OutputVariant]) -> list[dict[str, Any]]:
    rows: list[dict[str, Any]] = []
    for item in items:
        product = item.product
        variant = item.variant
        description_html = product.get("description", "")
        description = html_to_text(description_html)[:2000]
        bullets = extract_bullets(description_html)
        upper, sole, insole = infer_materials(product)
        price_divisor = item.source_pack_pairs if item.retail_from_pack else 1
        current_price = round(float(variant.get("price") or 0) / 100 / price_divisor, 2)
        compare_price_raw = variant.get("compare_at_price")
        compare_price = (
            round(float(compare_price_raw) / 100 / price_divisor, 2)
            if compare_price_raw
            else None
        )
        parent_sku = f"BELSTA-{product.get('id')}"
        child_sku = f"BELSTA-{variant.get('id')}"
        if item.retail_from_pack:
            child_sku += f"-{sku_size_token(item.size)}"
        product_title = normalize_space(product.get("title"))
        if item.retail_from_pack:
            product_title = retail_product_title(product_title)

        values: dict[str, Any] = {
            "E": item.category,
            "G": "Normal product",
            "L": product_title[:500],
            "M": parent_sku,
            "N": child_sku,
            "O": "Add",
            "T": description,
            "DS": upper,
            "DW": sole,
            "DY": "Fabric" if insole in {"Fabric", "Wool"} else insole,
            "EB": infer_closure(product),
            "FD": insole,
            "KF": "Color × Size",
            "KG": "3 - European Size",
            "KH": "7 - Numeric",
            "KI": item.size,
            "KJ": item.temu_color,
            "KK": "Add size chart manually",
            "KM": "cm-g-ml",
            "KN": item.size if item.size.isdigit() else "",
            "KO": foot_length(item.size, item.category),
            "LY": QUANTITY,
            "LZ": current_price,
            "MA": item.canonical_url,
            "MD": RETAIL_PAIR_WEIGHT_G if item.retail_from_pack else float(variant.get("weight") or 500),
            "ME": PACKAGE_LENGTH_CM,
            "MF": PACKAGE_WIDTH_CM,
            "MG": PACKAGE_HEIGHT_CM,
            "MH": "Single set",
            "MI": "Yes",
            "MJ": 1,
            "MK": "pair",
            "ML": 1,
            "MM": "pair",
            "MS": SHIPPING_TEMPLATE,
            "MT": HANDLING_TIME,
            "MU": FULFILLMENT_CHANNEL,
            "MW": COUNTRY_OF_ORIGIN,
        }

        if compare_price is not None and compare_price > current_price:
            values["MB"] = compare_price
        else:
            values["MC"] = "N/A"

        if item.category == CATEGORY_WOMEN_SLIPPERS:
            values.update({
                "EK": "Casual",
                "EZ": infer_pattern(product),
                "JY": "Standard Fit",
                "JZ": "Slide" if "отворен" in normalize_space(product.get("title")).casefold() else "Slip-on",
                "KA": "All-season",
                "KB": "Other Regions",
            })
        elif item.category == CATEGORY_MEN_SLIPPERS:
            values.update({
                "EF": "Round Toe",
                "EH": "Indoor",
                "EK": "Casual",
                "ET": infer_season(product),
                "EW": "Male",
                "EZ": infer_pattern(product),
                "JY": "Standard Fit",
            })
        elif item.category == CATEGORY_BOYS_SLIPPERS:
            values.update({
                "EF": "Round Toe",
                "EG": "14 And Under",
                "EH": "Indoor",
                "EK": "Casual",
                "EN": "Lightweight",
                "EQ": infer_popular_element(product),
                "ET": infer_season(product),
                "EW": "Boys",
                "EZ": infer_pattern(product),
                "FC": "Low top",
            })
        else:
            values.update({
                "EF": "Round Toe",
                "EG": "14 And Under",
                "EH": "Indoor",
                "EK": "Casual",
                "EN": "Lightweight",
                "EQ": infer_popular_element(product),
                "ET": infer_season(product),
                "EW": "Girls",
                "EZ": infer_pattern(product),
                "FC": "Low top",
            })

        for index, bullet in enumerate(bullets):
            values[chr(ord("U") + index)] = bullet

        # Detail image columns AA:DQ. Fill only the first 30.
        for index, image_url in enumerate(item.detail_images[:30]):
            values[column_number_to_letters(27 + index)] = image_url

        # SKU image columns LO:LX.
        start = column_index_from_string("LO")
        for index, image_url in enumerate(item.sku_images[:10]):
            values[column_number_to_letters(start + index)] = image_url

        rows.append({column: value for column, value in values.items() if value not in ("", None)})
    return rows


MAIN_NS = "http://schemas.openxmlformats.org/spreadsheetml/2006/main"
OFFICE_REL_NS = "http://schemas.openxmlformats.org/officeDocument/2006/relationships"
PACKAGE_REL_NS = "http://schemas.openxmlformats.org/package/2006/relationships"
XML_NS = "http://www.w3.org/XML/1998/namespace"
ET.register_namespace("", MAIN_NS)


def template_sheet_path(archive: zipfile.ZipFile) -> str:
    workbook_root = ET.fromstring(archive.read("xl/workbook.xml"))
    rels_root = ET.fromstring(archive.read("xl/_rels/workbook.xml.rels"))
    rel_targets = {
        rel.attrib["Id"]: rel.attrib["Target"]
        for rel in rels_root.findall(f"{{{PACKAGE_REL_NS}}}Relationship")
    }
    for sheet in workbook_root.findall(f".//{{{MAIN_NS}}}sheet"):
        if sheet.attrib.get("name") == "Template":
            rel_id = sheet.attrib.get(f"{{{OFFICE_REL_NS}}}id")
            target = rel_targets.get(rel_id or "")
            if not target:
                break
            return "xl/" + target.lstrip("/")
    raise RuntimeError("The workbook does not contain a Template sheet.")


def cell_column(cell_reference: str) -> int:
    match = re.match(r"([A-Z]+)", cell_reference)
    if not match:
        raise ValueError(f"Invalid cell reference: {cell_reference}")
    return column_index_from_string(match.group(1))


def set_xml_cell(row: ET.Element, row_number: int, column: str, value: Any) -> None:
    reference = f"{column}{row_number}"
    cell = next(
        (candidate for candidate in row.findall(f"{{{MAIN_NS}}}c") if candidate.attrib.get("r") == reference),
        None,
    )
    if cell is None:
        cell = ET.Element(f"{{{MAIN_NS}}}c", {"r": reference})
        row.append(cell)

    for child in list(cell):
        cell.remove(child)

    if isinstance(value, bool):
        cell.attrib["t"] = "b"
        ET.SubElement(cell, f"{{{MAIN_NS}}}v").text = "1" if value else "0"
    elif isinstance(value, (int, float)) and not isinstance(value, bool):
        cell.attrib.pop("t", None)
        number = f"{value:.10f}".rstrip("0").rstrip(".") if isinstance(value, float) else str(value)
        ET.SubElement(cell, f"{{{MAIN_NS}}}v").text = number
    else:
        cell.attrib["t"] = "inlineStr"
        inline = ET.SubElement(cell, f"{{{MAIN_NS}}}is")
        text_node = ET.SubElement(inline, f"{{{MAIN_NS}}}t")
        text_value = str(value)
        if text_value != text_value.strip() or "  " in text_value:
            text_node.attrib[f"{{{XML_NS}}}space"] = "preserve"
        text_node.text = text_value


def patch_template_xml(source_xml: bytes, rows_data: list[dict[str, Any]]) -> bytes:
    root = ET.fromstring(source_xml)
    sheet_data = root.find(f"{{{MAIN_NS}}}sheetData")
    if sheet_data is None:
        raise RuntimeError("Template sheet has no sheetData element.")
    rows_by_number = {
        int(row.attrib["r"]): row
        for row in sheet_data.findall(f"{{{MAIN_NS}}}row")
        if row.attrib.get("r", "").isdigit()
    }

    for row_number, values in enumerate(rows_data, start=5):
        row = rows_by_number.get(row_number)
        if row is None:
            row = ET.Element(f"{{{MAIN_NS}}}row", {"r": str(row_number)})
            sheet_data.append(row)
            rows_by_number[row_number] = row
        for column, value in values.items():
            set_xml_cell(row, row_number, column, value)
        cells = list(row.findall(f"{{{MAIN_NS}}}c"))
        cells.sort(key=lambda cell: cell_column(cell.attrib["r"]))
        for cell in cells:
            row.remove(cell)
        row.extend(cells)

    all_rows = list(sheet_data.findall(f"{{{MAIN_NS}}}row"))
    all_rows.sort(key=lambda row: int(row.attrib.get("r", "0")))
    for row in all_rows:
        sheet_data.remove(row)
    sheet_data.extend(all_rows)

    return ET.tostring(root, encoding="UTF-8", xml_declaration=True)


def write_workbook(template: Path, output_path: Path, items: list[OutputVariant]) -> None:
    if len(items) > 4996:
        raise ValueError("The Temu template supports at most 4,996 data rows.")
    output_path.parent.mkdir(parents=True, exist_ok=True)
    rows_data = build_workbook_rows(items)
    temporary_path = output_path.with_name(output_path.stem + ".tmp.xlsx")

    with zipfile.ZipFile(template, "r") as source_archive:
        sheet_path = template_sheet_path(source_archive)
        patched_sheet = patch_template_xml(source_archive.read(sheet_path), rows_data)
        with zipfile.ZipFile(temporary_path, "w") as target_archive:
            for info in source_archive.infolist():
                payload = patched_sheet if info.filename == sheet_path else source_archive.read(info.filename)
                target_archive.writestr(info, payload)

    temporary_path.replace(output_path)


def column_number_to_letters(number: int) -> str:
    letters = ""
    while number:
        number, remainder = divmod(number - 1, 26)
        letters = chr(65 + remainder) + letters
    return letters


def write_report(path: Path, items: list[OutputVariant]) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    with path.open("w", encoding="utf-8-sig", newline="") as stream:
        writer = csv.writer(stream)
        writer.writerow(
            [
                "Product",
                "Source color",
                "Temu color",
                "Size",
                "Source wholesale size",
                "Pairs in source pack",
                "Retail row from wholesale pack",
                "Source pack Base price EUR",
                "Retail Base price EUR",
                "Source pack List price EUR",
                "Retail List price EUR",
                "Shopify variant ID",
                "Source SKU",
                "Available",
                "Source weight g",
                "Retail weight g",
                "Category ID",
                "Reference link",
            ]
        )
        for item in items:
            variant = item.variant
            price_divisor = item.source_pack_pairs if item.retail_from_pack else 1
            source_base_price = float(variant.get("price") or 0) / 100
            source_list_price = (
                float(variant.get("compare_at_price") or 0) / 100
                if variant.get("compare_at_price")
                else ""
            )
            writer.writerow(
                [
                    normalize_space(item.product.get("title")),
                    item.source_color,
                    item.temu_color,
                    item.size,
                    item.source_size,
                    item.source_pack_pairs,
                    item.retail_from_pack,
                    source_base_price,
                    round(source_base_price / price_divisor, 2),
                    source_list_price,
                    round(float(source_list_price) / price_divisor, 2) if source_list_price != "" else "",
                    variant.get("id"),
                    normalize_space(variant.get("sku")),
                    variant.get("available"),
                    variant.get("weight"),
                    RETAIL_PAIR_WEIGHT_G if item.retail_from_pack else float(variant.get("weight") or 500),
                    item.category,
                    item.canonical_url,
                ]
            )


def write_summary(
    path: Path,
    selections: list[InputSelection],
    products: dict[str, dict[str, Any]],
    items: list[OutputVariant],
    warnings: list[str],
) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    exact_duplicates = len(selections) - len({item.original_url for item in selections})
    product_level_links = sum(item.variant_id is None for item in selections)
    selected_colors = {(item.product.get("id"), item.source_color) for item in items}
    source_variants = {(item.product.get("id"), item.variant.get("id")) for item in items}
    wholesale_source_variants = {
        (item.product.get("id"), item.variant.get("id"))
        for item in items
        if item.retail_from_pack
    }
    retail_rows_from_packs = sum(item.retail_from_pack for item in items)
    lines = [
        "BELSTA → Temu scraper summary",
        "",
        f"Input links: {len(selections)}",
        f"Exact duplicate links ignored: {exact_duplicates}",
        f"Product links without variant ID: {product_level_links}",
        f"Products fetched: {len(products)}",
        f"Selected product/color groups: {len(selected_colors)}",
        f"Source Shopify variants used: {len(source_variants)}",
        f"Wholesale source variants split by size: {len(wholesale_source_variants)}",
        f"Retail rows created from wholesale packs: {retail_rows_from_packs}",
        f"Output SKU rows: {len(items)}",
        f"Quantity per SKU: {QUANTITY}",
        f"Package dimensions: {PACKAGE_LENGTH_CM} × {PACKAGE_WIDTH_CM} × {PACKAGE_HEIGHT_CM} cm",
        "Shipping Template: blank",
        "Manufacturer: blank",
        "EU Responsible Person: blank",
        "",
        "Selection rule:",
        "- A link with a variant ID selects that variant's color and all currently available sizes in that color.",
        "- A product link without a variant ID selects the first available color shown by the product page and all sizes in that color.",
        "- A wholesale size range is expanded to one unique row for every individual size.",
        "- Wholesale pack prices are divided by the number of pairs in the source pack.",
        "- Duplicate product/color/size combinations are merged using the lowest per-pair source price.",
        "- Visually reviewed wholesale bundle images are excluded from SKU and detail images.",
        "",
    ]
    if warnings:
        lines.append("Warnings:")
        lines.extend(f"- {warning}" for warning in warnings)
    else:
        lines.append("Warnings: none")
    path.write_text("\n".join(lines) + "\n", encoding="utf-8")


def validate_output(path: Path, expected_rows: int) -> None:
    workbook = load_workbook(path, read_only=False, data_only=False)
    expected_sheets = {
        "Instructions",
        "Images",
        "Example",
        "Data Definitions",
        "Template",
        "Dropdown Lists",
        "GoodsLevelMode",
        "Browse Data",
        "Valid Values",
        "Conditions List",
        "Category Name",
    }
    if set(workbook.sheetnames) != expected_sheets:
        raise RuntimeError("The output workbook sheet structure differs from the Temu template.")
    sheet = workbook["Template"]
    populated = sum(1 for row in range(5, 5 + expected_rows) if sheet.cell(row, 13).value)
    if populated != expected_rows:
        raise RuntimeError(f"Expected {expected_rows} data rows but found {populated}.")
    dropdown_sheet = workbook["Dropdown Lists"]
    color_options: dict[str, set[str]] = {}
    for row_values in dropdown_sheet.iter_rows(values_only=True):
        key = str(row_values[0] or "")
        match = re.fullmatch(r"t_4_(\d+)_Color", key)
        if match:
            color_options[match.group(1)] = {
                str(value) for value in row_values[2:] if value is not None
            }

    contribution_skus: set[str] = set()
    variation_combinations: set[tuple[str, str, str]] = set()
    for row in range(5, 5 + expected_rows):
        required = {
            "Category": sheet.cell(row, column_index_from_string("E")).value,
            "Product Name": sheet.cell(row, column_index_from_string("L")).value,
            "Contribution Goods": sheet.cell(row, column_index_from_string("M")).value,
            "Contribution SKU": sheet.cell(row, column_index_from_string("N")).value,
            "SKU image": sheet.cell(row, column_index_from_string("LO")).value,
            "Quantity": sheet.cell(row, column_index_from_string("LY")).value,
            "Base Price": sheet.cell(row, column_index_from_string("LZ")).value,
            "Weight": sheet.cell(row, column_index_from_string("MD")).value,
        }
        missing = [name for name, value in required.items() if value in (None, "")]
        if missing:
            raise RuntimeError(f"Output row {row} is missing: {', '.join(missing)}")
        category = str(required["Category"])
        color = str(sheet.cell(row, column_index_from_string("KJ")).value or "")
        if color not in color_options.get(category, set()):
            raise RuntimeError(f"Output row {row} has invalid Temu color {color!r} for category {category}.")
        contribution_sku = str(required["Contribution SKU"])
        if contribution_sku in contribution_skus:
            raise RuntimeError(f"Duplicate Contribution SKU in output row {row}: {contribution_sku}")
        contribution_skus.add(contribution_sku)
        variation_key = (
            str(required["Contribution Goods"]),
            color,
            str(sheet.cell(row, column_index_from_string("KI")).value or ""),
        )
        if variation_key in variation_combinations:
            raise RuntimeError(
                f"Duplicate product/color/size combination in output row {row}: {variation_key}"
            )
        variation_combinations.add(variation_key)
        if sheet.cell(row, column_index_from_string("LY")).value != QUANTITY:
            raise RuntimeError(f"Output row {row} has an unexpected quantity.")
        dimensions = tuple(
            sheet.cell(row, column_index_from_string(column)).value for column in ("ME", "MF", "MG")
        )
        if dimensions != (PACKAGE_LENGTH_CM, PACKAGE_WIDTH_CM, PACKAGE_HEIGHT_CM):
            raise RuntimeError(f"Output row {row} has unexpected package dimensions: {dimensions}")
        image_columns = list(range(column_index_from_string("AA"), column_index_from_string("DQ") + 1))
        image_columns.extend(
            range(column_index_from_string("LO"), column_index_from_string("LX") + 1)
        )
        for column in image_columns:
            image_url = sheet.cell(row, column).value
            if is_wholesale_image_url(image_url):
                raise RuntimeError(
                    f"Output row {row} contains a wholesale bundle image in column {column}."
                )
        base_price = float(sheet.cell(row, column_index_from_string("LZ")).value)
        list_price = sheet.cell(row, column_index_from_string("MB")).value
        no_list_price = sheet.cell(row, column_index_from_string("MC")).value
        if list_price not in (None, "") and float(list_price) <= base_price:
            raise RuntimeError(f"Output row {row} has List Price not greater than Base Price.")
        if list_price in (None, "") and no_list_price != "N/A":
            raise RuntimeError(f"Output row {row} has neither a valid List Price nor N/A.")
        for cell in sheet[row]:
            if isinstance(cell.value, str) and cell.value.startswith("#"):
                raise RuntimeError(f"Formula error marker {cell.value!r} in {cell.coordinate}.")


def main() -> int:
    args = parse_args()
    logging.basicConfig(level=logging.INFO, format="%(levelname)s: %(message)s")
    available_path = Path(args.available)
    template_path = Path(args.template)
    output_path = Path(args.output)
    report_path = Path(args.report)
    summary_path = Path(args.summary)

    for required_path in (available_path, template_path):
        if not required_path.exists():
            raise FileNotFoundError(f"Missing input file: {required_path}")

    selections = read_selections(available_path)
    handles = sorted({selection.handle for selection in selections})
    logging.info("Input contains %d links and %d unique products", len(selections), len(handles))
    products = fetch_products(handles)
    items, warnings = select_variants(selections, products)
    if not items:
        raise RuntimeError("No available variants matched the input links.")

    items.sort(
        key=lambda item: (
            normalize_space(item.product.get("title")).casefold(),
            item.temu_color.casefold(),
            natural_size_key(item.size),
            int(item.variant.get("id") or 0),
        )
    )
    logging.info("Writing %d selected SKU rows", len(items))
    write_workbook(template_path, output_path, items)
    write_report(report_path, items)
    write_summary(summary_path, selections, products, items, warnings)
    validate_output(output_path, len(items))
    logging.info("Completed successfully: %s", output_path)
    return 0


def natural_size_key(value: str) -> tuple[Any, ...]:
    parts = re.split(r"(\d+(?:[.,]\d+)?)", normalize_space(value))
    key: list[Any] = []
    for part in parts:
        if re.fullmatch(r"\d+(?:[.,]\d+)?", part):
            key.append((0, float(part.replace(",", "."))))
        else:
            key.append((1, part.casefold()))
    return tuple(key)


if __name__ == "__main__":
    try:
        raise SystemExit(main())
    except Exception as exc:  # noqa: BLE001 - print a clear Actions error.
        logging.exception("Scraper failed: %s", exc)
        raise SystemExit(1) from exc
